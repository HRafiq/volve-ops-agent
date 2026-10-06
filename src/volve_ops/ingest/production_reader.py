"""Read the Volve production workbook into canonical records.

This is the only module that knows the source's column names. Everything downstream sees
`ProductionDay` and nothing else, so a change in the source shape is a change here and
nowhere else.

openpyxl rather than pandas, deliberately. An empty numeric cell comes back from openpyxl as
None and from pandas as NaN. None is what this project means by missing; NaN would pass the
"is a float" test, fail every physical check silently, land in no day class at all, and then
propagate through every sum downstream.
"""

from __future__ import annotations

import datetime as dt
from collections.abc import Iterator
from pathlib import Path
from typing import Final

from openpyxl import load_workbook

from volve_ops import INGEST_VERSION
from volve_ops.ingest.production import ProductionDay, WellStatus
from volve_ops.ingest.quarantine import UnsafeInputError

DAILY_SHEET: Final[str] = "Daily Production Data"

# Source column -> canonical field. Volumes in this workbook are already Sm3 and on-stream
# time is already hours, so no conversion is applied; units.py exists for sources that need it.
COLUMNS: Final[dict[str, str]] = {
    "DATEPRD": "production_date",
    "NPD_WELL_BORE_NAME": "well",
    "ON_STREAM_HRS": "on_stream_hours",
    "BORE_OIL_VOL": "oil_volume_sm3",
    "BORE_GAS_VOL": "gas_volume_sm3",
    "BORE_WAT_VOL": "water_volume_sm3",
    "AVG_CHOKE_SIZE_P": "choke_size",
    "AVG_CHOKE_UOM": "choke_unit",
    "AVG_DOWNHOLE_PRESSURE": "downhole_pressure_bar",
    "AVG_DOWNHOLE_TEMPERATURE": "downhole_temperature_c",
    "AVG_DP_TUBING": "tubing_dp_bar",
    "AVG_ANNULUS_PRESS": "annulus_pressure_bar",
    "AVG_WHP_P": "wellhead_pressure_bar",
    "AVG_WHT_P": "wellhead_temperature_c",
    "DP_CHOKE_SIZE": "choke_dp_bar",
}

# The workbook carries both FLOW_KIND and WELL_TYPE. They disagree on a handful of rows, where
# a wellbore typed as an injector reports a production flow. FLOW_KIND wins, because it
# describes what the well did on that day while WELL_TYPE describes how it is classified
# overall, and the day classification in docs/eval_protocol.md section 6 is about the day.
FLOW_KIND_TO_STATUS: Final[dict[str, WellStatus]] = {
    "production": WellStatus.PRODUCING,
    "injection": WellStatus.INJECTING,
}


class SourceSchemaError(UnsafeInputError):
    """The workbook does not have the columns this reader was written against."""


def _as_float(value: object, column: str, row: int) -> float | None:
    """Coerce a cell to a float or None, refusing anything that is neither."""
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise SourceSchemaError(f"row {row}: {column} holds a boolean, not a number")
    if isinstance(value, int | float):
        return float(value)
    raise SourceSchemaError(f"row {row}: {column} holds {value!r}, which is not a number")


def _as_date(value: object, row: int) -> dt.date:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    raise SourceSchemaError(f"row {row}: DATEPRD holds {value!r}, which is not a date")


def read_daily_production(path: Path, *, source: str | None = None) -> Iterator[ProductionDay]:
    """Yield one ProductionDay per row of the daily sheet.

    Raises SourceSchemaError if a required column is missing or a cell holds something the
    canonical model cannot represent. A source that has changed shape is a problem to surface,
    not one to work around row by row.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if DAILY_SHEET not in workbook.sheetnames:
            raise SourceSchemaError(f"no sheet named {DAILY_SHEET!r}; found {workbook.sheetnames}")
        sheet = workbook[DAILY_SHEET]
        rows = sheet.iter_rows(values_only=True)

        try:
            header = [str(c).strip() if c is not None else "" for c in next(rows)]
        except StopIteration:
            raise SourceSchemaError("the daily sheet is empty") from None

        index = {name: i for i, name in enumerate(header)}
        missing = [c for c in (*COLUMNS, "FLOW_KIND") if c not in index]
        if missing:
            raise SourceSchemaError(f"missing expected columns: {missing}")

        label = source or f"{path.name}:{DAILY_SHEET}"

        for number, row in enumerate(rows, start=2):
            if all(cell is None for cell in row):
                continue

            flow_kind = row[index["FLOW_KIND"]]
            if flow_kind is not None and str(flow_kind).strip() not in FLOW_KIND_TO_STATUS:
                raise SourceSchemaError(
                    f"row {number}: unmapped FLOW_KIND {flow_kind!r}; "
                    f"known values {sorted(FLOW_KIND_TO_STATUS)}"
                )
            status = FLOW_KIND_TO_STATUS[str(flow_kind).strip()] if flow_kind is not None else None

            well = row[index["NPD_WELL_BORE_NAME"]]
            if well is None or not str(well).strip():
                raise SourceSchemaError(f"row {number}: no well name")

            choke_unit = row[index["AVG_CHOKE_UOM"]]

            yield ProductionDay(
                well=str(well).strip(),
                production_date=_as_date(row[index["DATEPRD"]], number),
                on_stream_hours=_as_float(row[index["ON_STREAM_HRS"]], "ON_STREAM_HRS", number),
                oil_volume_sm3=_as_float(row[index["BORE_OIL_VOL"]], "BORE_OIL_VOL", number),
                gas_volume_sm3=_as_float(row[index["BORE_GAS_VOL"]], "BORE_GAS_VOL", number),
                water_volume_sm3=_as_float(row[index["BORE_WAT_VOL"]], "BORE_WAT_VOL", number),
                well_status=status,
                choke_size=_as_float(row[index["AVG_CHOKE_SIZE_P"]], "AVG_CHOKE_SIZE_P", number),
                choke_unit=str(choke_unit).strip() if choke_unit is not None else None,
                downhole_pressure_bar=_as_float(
                    row[index["AVG_DOWNHOLE_PRESSURE"]], "AVG_DOWNHOLE_PRESSURE", number
                ),
                downhole_temperature_c=_as_float(
                    row[index["AVG_DOWNHOLE_TEMPERATURE"]], "AVG_DOWNHOLE_TEMPERATURE", number
                ),
                tubing_dp_bar=_as_float(row[index["AVG_DP_TUBING"]], "AVG_DP_TUBING", number),
                annulus_pressure_bar=_as_float(
                    row[index["AVG_ANNULUS_PRESS"]], "AVG_ANNULUS_PRESS", number
                ),
                wellhead_pressure_bar=_as_float(row[index["AVG_WHP_P"]], "AVG_WHP_P", number),
                wellhead_temperature_c=_as_float(row[index["AVG_WHT_P"]], "AVG_WHT_P", number),
                choke_dp_bar=_as_float(row[index["DP_CHOKE_SIZE"]], "DP_CHOKE_SIZE", number),
                source=label,
                parser_version=INGEST_VERSION,
            )
    finally:
        workbook.close()
